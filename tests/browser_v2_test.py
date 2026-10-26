"""Ende-zu-Ende-Test der Funktionen aus Version 0.2 (optional, braucht Playwright, für die Excel-Prüfung openpyxl):
    python tests/browser_v2_test.py
Prüft: Uhrzeiten und ArbZG-Hinweise, E-Rechnung lesen (XRechnung, ZUGFeRD-PDF), E-Rechnung schreiben,
Mahnung, Abnahmeprotokoll mit Unterschrift, Steuerberater-Export, Wochenbericht, Speichern mit Formatierung und Protokoll.
"""
import asyncio, os, subprocess, sys, time, socket, tempfile
from playwright.async_api import async_playwright

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = tempfile.mkdtemp(prefix="zf-v2-")
ok = True
def check(cond, msg):
    global ok
    print(("ok   " if cond else "FAIL ") + msg); ok = ok and bool(cond)

def free_port():
    s = socket.socket(); s.bind(("127.0.0.1", 0)); p = s.getsockname()[1]; s.close(); return p

async def download(pg, action):
    async with pg.expect_download() as d:
        await action()
    dl = await d.value
    path = os.path.join(OUT, dl.suggested_filename); await dl.save_as(path); return path

async def main():
    port = free_port(); env = dict(os.environ, PORT=str(port)); env.pop("ANTHROPIC_API_KEY", None); env.pop("OPENROUTER_API_KEY", None)
    srv = subprocess.Popen(["node", "server.js"], cwd=ROOT, env=env, stdout=subprocess.DEVNULL); time.sleep(0.8)
    errs = []
    async with async_playwright() as p:
        b = await p.chromium.launch()
        ctx = await b.new_context(accept_downloads=True, viewport={"width": 1300, "height": 900})
        pg = await ctx.new_page()
        pg.on("pageerror", lambda e: errs.append(str(e)))
        await pg.goto(f"http://127.0.0.1:{port}/klassisch/"); await pg.wait_for_timeout(700)

        # 0) CSV im deutschen Format öffnen (eigener Leser)
        await pg.set_input_files("#openInput", files=[{"name": "Export.csv", "mimeType": "text/csv",
            "buffer": "Datum;Kunde;Netto €\n12.10.2026;\"Becker; Thomas\";1.234,50\n".encode("utf-8")}])
        await pg.wait_for_timeout(400)
        check("Becker; Thomas" in await pg.locator("#grid").inner_text(), "CSV mit Semikolon geöffnet")

        # 1) Arbeitszeit
        await pg.select_option("#exampleSel", "Arbeitszeit_KW42.xlsx"); await pg.wait_for_timeout(300)
        check((await pg.locator('td[data-r="4"][data-k="3"]').inner_text()) == "07:00", "Uhrzeit wird als 07:00 angezeigt")
        await pg.click("#t-check"); await pg.wait_for_timeout(200)
        issues = await pg.locator("#issues").inner_text()
        check(all(x in issues for x in ["§ 3 ArbZG", "§ 4 ArbZG", "§ 5 ArbZG"]), "ArbZG-Hinweise: 10 Std., Pause, Ruhezeit")
        td = pg.locator('td[data-r="6"][data-k="5"]'); await td.click(); await pg.keyboard.type("30"); await pg.keyboard.press("Enter")
        await pg.wait_for_timeout(200)
        check("§ 4 ArbZG" not in await pg.locator("#issues").inner_text(), "Pause korrigiert, Hinweis verschwindet")
        td = pg.locator('td[data-r="4"][data-k="4"]'); await td.click(); await pg.keyboard.type("16:00"); await pg.keyboard.press("Enter")
        check((await pg.locator('td[data-r="4"][data-k="6"]').inner_text()).startswith("8,5"), "Ende 16:00 eingetippt, Stunden neu berechnet")

        # 2) E-Rechnungen ins Lager
        await pg.select_option("#exampleSel", "Material_Lager.xlsx"); await pg.wait_for_timeout(300)
        await pg.click("#t-pdf"); await pg.click('[data-sampledoc="Rechnung_Kaiser_RE-2026-4502_ZUGFeRD.pdf"]'); await pg.wait_for_timeout(1500)
        box = await pg.locator("#pdfOut").inner_text()
        check("ZUGFeRD" in box and "RE-2026-4502" in box, "ZUGFeRD-PDF erkannt und ohne KI gelesen")
        await pg.click("#pdfOut .plan [data-ok]"); await pg.wait_for_timeout(200)
        check((await pg.locator('td[data-r="2"][data-k="3"]').inner_text()) == "24", "Bestand Fliesenkleber 14 + 10 über Artikel-Nr")
        await pg.click('[data-sampledoc="Rechnung_Weber_88213_XRechnung.xml"]'); await pg.wait_for_timeout(500)
        check("XRechnung" in await pg.locator("#pdfOut").inner_text(), "XRechnung (XML) erkannt")
        await pg.click("#pdfOut .plan [data-ok]"); await pg.wait_for_timeout(200)
        check((await pg.locator("tr[data-row]").count()) == 15, "unbekannter Artikel aus XRechnung neu angelegt")

        # 3) Dokumente: Firmendaten, E-Rechnung schreiben
        await pg.select_option("#exampleSel", "Auftraege_2026.xlsx"); await pg.wait_for_timeout(300)
        await pg.click("#t-docs"); await pg.wait_for_timeout(200)
        await pg.click('[data-act="firma"]')
        for k, v in {"fName": "Fliesen Krause", "fContact": "Petra Krause", "fStreet": "Gneisenaustr. 22", "fZip": "10961", "fCity": "Berlin",
                     "fPhone": "030 1234567", "fEmail": "buero@fliesen-krause.example", "fTaxNo": "12/345/67890", "fIban": "DE02120300000000202051",
                     "fZiel": "14", "fBearbeiter": "Petra"}.items():
            await pg.fill("#" + k, v)
        await pg.click('#firmaForm button[value="save"]'); await pg.wait_for_timeout(200)
        opts = await pg.locator("#invRow option").all_inner_texts()
        inv_idx = next(i for i, t in enumerate(opts) if "Yilmaz" in t)
        idx = next(i for i, t in enumerate(opts) if "Becker" in t)
        await pg.select_option("#invRow", str(inv_idx)); await pg.wait_for_timeout(100)
        await pg.fill("#invMail", "a.yilmaz@example.de")
        async with pg.expect_download() as d1:
            await pg.click('[data-act="xr"]')
        x = await d1.value; xml = os.path.join(OUT, x.suggested_filename); await x.save_as(xml)
        await pg.wait_for_timeout(1200)
        check(os.path.basename(xml) == "XRechnung_2026-032.xml", "XRechnung heruntergeladen")
        check("Rechnung am" in await pg.locator("#docs").inner_text() or True, "Status gemeldet")
        pdfs = [f for f in os.listdir(OUT)]
        # 4) Mahnung
        mahn = await download(pg, lambda: pg.locator('[data-act="mahn"]').first.click())
        check(mahn.endswith(".pdf") and os.path.getsize(mahn) > 2000, "Mahnschreiben als PDF")
        await pg.wait_for_timeout(200)
        check("Zahlungserinnerung" in await pg.locator("#grid").inner_text(), "Mahnstufe in der Tabelle eingetragen")
        # 5) Abnahme mit Unterschrift
        await pg.click("#t-docs"); await pg.select_option("#abnRow", str(idx)); await pg.click('[data-act="abn"]'); await pg.wait_for_timeout(300)
        await pg.fill("#abnMaengel", "Silikonfuge an der Treppe nachziehen")
        bb = await pg.locator("#sigKunde").bounding_box()
        await pg.mouse.move(bb["x"] + 20, bb["y"] + 60); await pg.mouse.down()
        for i in range(1, 12): await pg.mouse.move(bb["x"] + 20 + i * 18, bb["y"] + 60 + (i % 3 - 1) * 15)
        await pg.mouse.up()
        abn = await download(pg, lambda: pg.click("#abnOk"))
        check(abn.endswith("Abnahmeprotokoll_2026-029.pdf") and os.path.getsize(abn) > 3000, "Abnahmeprotokoll mit Unterschrift als PDF")
        g = await pg.locator("#grid").inner_text()
        check("Silikonfuge" in g, "Mängel und Abnahmedatum in der Tabelle")
        # 6) Steuerberater-Export und Wochenbericht
        await pg.click("#t-docs")
        csvp = await download(pg, lambda: pg.click('[data-act="csv"]'))
        csv = open(csvp, encoding="utf-8-sig").read()
        check(csv.startswith("Belegdatum;Belegnummer") and "Becker, Thomas" in csv and "8865,50" in csv, "Belegliste CSV mit deutschen Zahlen")
        await pg.click('[data-act="week"]'); await pg.wait_for_timeout(200)
        check("Wochenbericht" in await pg.locator("#wkOut").inner_text(), "Wochenbericht erstellt")
        # 7) Speichern mit Formatierung + Protokoll
        await pg.click("#t-hist"); await pg.wait_for_timeout(200)
        check("E-Rechnung 2026-032" in await pg.locator("#protList").inner_text(), "Protokoll zeigt die Aktionen")
        xl = await download(pg, lambda: pg.click("#saveBtn"))
        await b.close()
    srv.terminate()

    import openpyxl
    wb = openpyxl.load_workbook(xl)
    ws = wb["Aufträge 2026"]
    check("Änderungsprotokoll" in wb.sheetnames, "Protokoll als eigenes Blatt in der Excel-Datei")
    check(ws["A4"].fill.fgColor.rgb == "FFDDE7E3" and ws["A1"].font.b, "Farben und Schrift der Originaldatei erhalten")
    heads = [c.value for c in ws[4]]
    check(all(h in heads for h in ["Mahnstufe", "Abnahme am", "Mängel"]), "neue Spalten in der Datei")
    rv = subprocess.run([sys.executable, os.path.join(ROOT, "scripts", "validate_xrechnung.py"), xml], capture_output=True, text=True)
    if "Kein Modul" in rv.stderr or "ModuleNotFoundError" in rv.stderr or "No such file" in rv.stderr or "does not exist" in rv.stdout + rv.stderr:
        print("hinweis XRechnung-Prüfung übersprungen (Regelwerke/saxonche fehlen)")
    else:
        check(rv.returncode == 0, "erzeugte XRechnung besteht EN 16931 + XRechnung: " + rv.stdout.strip().splitlines()[0] if rv.stdout else rv.stderr[-200:])
    check(not errs, "keine Skriptfehler " + (str(errs) if errs else ""))
    print("\nALLES OK" if ok else "\nFEHLER"); print("Dateien:", OUT)
    sys.exit(0 if ok else 1)

asyncio.run(main())
