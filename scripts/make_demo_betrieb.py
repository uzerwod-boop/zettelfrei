"""Große Demo-Datei „Krause_Betrieb_Oktober_2026.xlsx“: so sieht ein echtes Betriebsbuch aus.
Zweizeilige Überschriften, Wochen-Zwischensummen, Nebentabelle mit SVERWEIS, verstecktes Kundenblatt,
Übersicht mit Bezügen auf andere Blätter – und typische Fehler zum Finden.
    python scripts/make_demo_betrieb.py   (openpyxl; LibreOffice optional für zwischengespeicherte Werte)
"""
import datetime as dt, os, shutil, subprocess, tempfile
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "beispiele", "Krause_Betrieb_Oktober_2026.xlsx")
B = Font(bold=True); W = Font(bold=True, color="FFFFFF")
HEAD = PatternFill("solid", fgColor="0F6B5C"); SUB = PatternFill("solid", fgColor="DDE7E3"); WEEK = PatternFill("solid", fgColor="FFF2CC")
TOT = PatternFill("solid", fgColor="F6DD6A"); thin = Side(style="thin", color="B7C2BE"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
EUR = '#,##0.00 "€"'; DATE = "DD.MM.YYYY"; TIME = "hh:mm"
def T(h, m=0): return dt.time(h, m)
wb = Workbook()

# ---------------- Übersicht ----------------
u = wb.active; u.title = "Übersicht"
u["B2"] = "Fliesen Krause – Monatsübersicht Oktober 2026"; u["B2"].font = Font(bold=True, size=16)
u["B3"] = "automatisch – hier bitte nichts eintragen"; u["B3"].font = Font(italic=True, color="7F7F7F")

# ---------------- Stunden ----------------
s = wb.create_sheet("Stunden Oktober")
s["B1"] = "Fliesen Krause GmbH · Gneisenaustr. 22 · Berlin"; s["B1"].font = Font(bold=True, size=14)
s["B2"] = "Stundenerfassung Oktober 2026 – bitte jeden Freitag ausfüllen!"; s["B2"].font = Font(color="C00000", bold=True)
heads = [("B4:B5", "Datum", None), ("C4:D4", "Wer / Wo", ["Mitarbeiter", "Baustelle"]), ("E4:G4", "Arbeitszeit", ["Beginn", "Ende", "Pause (Min.)"]),
         ("H4:H5", "Stunden", None), ("I4:I5", "Lohn €", None), ("J4:J5", "Tätigkeit / Notiz", None)]
for rng, title, subs in heads:
    s.merge_cells(rng); a = s[rng.split(":")[0]]; a.value = title
    if subs:
        c0 = a.column
        for i, t in enumerate(subs): s.cell(5, c0 + i, t)
for r in (4, 5):
    for c in range(2, 11):
        x = s.cell(r, c); x.font = W; x.fill = HEAD; x.border = BOX; x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
# Nebentabelle: Stundensätze
s["L4"] = "Mitarbeiter"; s["M4"] = "Satz €/Std."
for c in ("L4", "M4"): s[c].font = W; s[c].fill = HEAD
RATES = [("Mehmet", 52), ("Kevin", 45), ("Jonas", 45), ("Ayşe", 38)]
for i, (n, v) in enumerate(RATES): s.cell(5 + i, 12, n); s.cell(5 + i, 13, v).number_format = EUR
s["L10"] = "Sätze ab 01.07.2026"; s["L10"].font = Font(italic=True, color="7F7F7F")

weeks = [
    ("KW 40", [(dt.date(2026,10,1), "Mehmet", "Bad Schneider", T(7), T(16), 30, "Wand gefliest"),
               (dt.date(2026,10,1), "Kevin", "Bad Schneider", T(7), T(15,30), 30, "Estrich geprüft"),
               (dt.date(2026,10,2), "Mehmet", "Bad Schneider", T(7), T(15), 30, "Fugen"),
               (dt.date(2026,10,2), "Ayşe", "Büro / Lager", T(8), T(12,30), 0, "Bestellung Kaiser")]),
    ("KW 41", [(dt.date(2026,10,5), "Jonas", "Praxis Dr. Hahn", T(7), T(16,30), 45, "Boden 60x60"),
               ("06.10.26", "Mehmet", "Küche Yılmaz", "7:00", "16:00", "30 min", "Fliesenspiegel"),
               (dt.date(2026,10,6), "Kevin", "Praxis Dr. Hahn", T(7), T(17), 30, "Sockel"),
               (dt.date(2026,10,7), "Jonas", "Praxis Dr. Hahn", T(7), T(15,30), 30, ""),
               (dt.date(2026,10,8), "Kevin", "Küche Yılmaz", T(6,30), T(15,30), 15, "Silikon"),
               (dt.date(2026,10,9), "Ayşe", "Büro / Lager", T(8), T(13), 0, "Inventur")]),
    ("KW 42", [(dt.date(2026,10,12), "Jonas", "Neubau Weber", T(7), T(18,30), 30, "Estrich noch feucht – Überstunden"),
               (dt.date(2026,10,13), "Jonas", "Neubau Weber", T(5), T(13), 30, "früh wegen Lieferung"),
               (dt.date(2026,10,13), "Mehmet", "Neubau Weber", T(7), T(16), 30, "Bad OG"),
               (dt.date(2026,10,14), "mehmet ", "Neubau Weber", T(7), T(16), 30, "Bad OG"),
               (dt.date(2026,10,15), "Kevin", "Bad Schneider", T(7), T(12), 0, "Nacharbeit Silikon"),
               (dt.date(2026,10,16), "Mehmet", "Neubau Weber", T(7), T(14), 30, "Treppe")]),
    ("KW 43", [(dt.date(2026,10,19), "Mehmet", "Neubau Weber", T(7), T(16), 30, "Bad EG"),
               (dt.date(2026,10,19), "Kevin", "Neubau Weber", T(7), T(16), 30, "Bad EG")]),
]
r = 6; first = r; subs = []
for wk, rows in weeks:
    s.cell(r, 2, f"— {wk} —").font = Font(bold=True, italic=True, color="0F6B5C"); r += 1
    ws0 = r
    for d, ma, bs, b, e, p, note in rows:
        s.cell(r, 2, d); s.cell(r, 3, ma); s.cell(r, 4, bs); s.cell(r, 5, b); s.cell(r, 6, e); s.cell(r, 7, p); s.cell(r, 10, note)
        if isinstance(d, dt.date): s.cell(r, 2).number_format = DATE
        for c in (5, 6):
            if isinstance(s.cell(r, c).value, dt.time): s.cell(r, c).number_format = TIME
        s.cell(r, 8, f'=IFERROR(ROUND((F{r}-E{r})*24-IF(ISNUMBER(G{r}),G{r},0)/60,2),"")'); s.cell(r, 8).number_format = "0.00"
        s.cell(r, 9, f'=IF(ISNUMBER(H{r}),ROUND(H{r}*VLOOKUP(TRIM(C{r}),$L$5:$M$8,2,FALSE),2),"")'); s.cell(r, 9).number_format = EUR
        for c in range(2, 11): s.cell(r, c).border = BOX
        r += 1
    if wk != "KW 43":
        s.cell(r, 2, f"Zwischensumme {wk}").font = B
        s.cell(r, 8, f"=SUBTOTAL(9,H{ws0}:H{r-1})"); s.cell(r, 9, f"=SUBTOTAL(9,I{ws0}:I{r-1})")
        s.cell(r, 8).number_format = "0.00"; s.cell(r, 9).number_format = EUR
        for c in range(2, 11): s.cell(r, c).fill = WEEK; s.cell(r, c).border = BOX
        r += 1
last = r - 1
s.cell(r, 2, "GESAMT Oktober").font = B
s.cell(r, 8, f"=SUBTOTAL(9,H{first}:H{last})"); s.cell(r, 9, f"=SUBTOTAL(9,I{first}:I{last})")
s.cell(r, 8).number_format = "0.00"; s.cell(r, 9).number_format = EUR
for c in range(2, 11): s.cell(r, c).fill = TOT; s.cell(r, c).font = B; s.cell(r, c).border = BOX
STD_TOTAL = r
s.cell(r + 2, 2, "Stand:"); s.cell(r + 2, 3, dt.date(2026, 10, 19)).number_format = DATE
s.conditional_formatting.add(f"H{first}:H{last}", CellIsRule(operator="greaterThan", formula=["10"], fill=PatternFill("solid", fgColor="F4CCCC")))
s.freeze_panes = "C6"
for c, w in zip("ABCDEFGHIJKLM", (2, 12, 13, 18, 9, 9, 9, 9, 11, 30, 3, 13, 11)): s.column_dimensions[c].width = w

# ---------------- Aufträge ----------------
kd = wb.create_sheet("Kunden"); kd.sheet_state = "hidden"
kd.append(["Kd-Nr", "Name", "Zahlungsziel", "E-Mail"])
for row in [(101, "Becker, Thomas", 14, "t.becker@example.de"), (102, "Yılmaz GbR", 30, "info@yilmaz.example"),
            (103, "Praxis Dr. Hahn", 10, "praxis@hahn.example"), (104, "Weber Bau GmbH", 30, "rechnung@weber-bau.example"),
            (105, "Schneider, Anna", 14, "a.schneider@example.de"), (106, "Hausverwaltung Lindner", 21, "buchhaltung@lindner.example")]:
    kd.append(row)
a = wb.create_sheet("Aufträge")
a["A1"] = "Aufträge und Rechnungen 2026"; a["A1"].font = Font(bold=True, size=14)
H = ["Auftrags-Nr", "Kd-Nr", "Kunde", "Baustelle", "Netto €", "MwSt-Satz", "MwSt €", "Brutto €", "Rechnung am", "Fällig am", "Status", "bezahlt am"]
for i, h in enumerate(H, 1): x = a.cell(3, i, h); x.font = W; x.fill = HEAD; x.border = BOX; x.alignment = Alignment(horizontal="center", wrap_text=True)
orders = [("2026-028", 105, "Bad Schneider", 6420, 0.19, dt.date(2026,9,10), "bezahlt", dt.date(2026,9,22)),
          ("2026-029", 101, "Gäste-WC Becker", 2380, 0.19, dt.date(2026,9,25), "offen", None),
          ("2026-030", 103, "Praxis Dr. Hahn", 9850, 0.19, dt.date(2026,10,2), "offen", None),
          ("2026-031", 102, "Küche Yılmaz", 3120, 0.19, dt.date(2026,10,9), "Teilzahlung", None),
          ("2026-031", 106, "Treppenhaus Lindner", "4.250,00", 0.19, dt.date(2026,10,12), "offen", None),
          ("2026-033", 104, "Neubau Weber (Abschlag 1)", 15000, 19, dt.date(2026,10,14), "offen", None),
          ("2026-034", 107, "Balkon Krüger", 1890, 0.07, dt.date(2026,10,15), "offen", None),
          ("2026-035", 104, "Neubau Weber (Abschlag 2)", 12000, 0.19, None, "Angebot", None)]
for i, (nr, k, bs, net, st, rd, stat, paid) in enumerate(orders):
    rr = 4 + i
    a.cell(rr, 1, nr); a.cell(rr, 2, k)
    a.cell(rr, 3, f'=IFERROR(VLOOKUP(B{rr},Kunden!$A:$D,2,FALSE),"??? unbekannt")')
    a.cell(rr, 4, bs); a.cell(rr, 5, net).number_format = EUR; a.cell(rr, 6, st).number_format = "0%"
    a.cell(rr, 7, f"=IFERROR(E{rr}*F{rr},0)").number_format = EUR; a.cell(rr, 8, f"=IFERROR(E{rr}+G{rr},E{rr})").number_format = EUR
    if rd: a.cell(rr, 9, rd).number_format = DATE
    a.cell(rr, 10, f'=IF(I{rr}="","",I{rr}+IFERROR(VLOOKUP(B{rr},Kunden!$A:$C,3,FALSE),14))').number_format = DATE
    a.cell(rr, 11, stat); 
    if paid: a.cell(rr, 12, paid).number_format = DATE
    for c in range(1, 13): a.cell(rr, c).border = BOX
lr = 4 + len(orders) - 1
a.cell(lr + 2, 4, "Summe").font = B
for c in (5, 7, 8): a.cell(lr + 2, c, f"=SUM({chr(64+c)}4:{chr(64+c)}{lr})").number_format = EUR; a.cell(lr + 2, c).font = B
a.cell(lr + 3, 4, "davon offen:").font = B
a.cell(lr + 3, 8, f'=SUMIF(K4:K{lr},"offen",H4:H{lr})').number_format = EUR; a.cell(lr + 3, 8).font = B
dv = DataValidation(type="list", formula1='"Angebot,offen,Teilzahlung,bezahlt,Mahnung"'); a.add_data_validation(dv); dv.add(f"K4:K60")
for c, w in zip("ABCDEFGHIJKL", (12, 7, 22, 24, 12, 9, 11, 12, 12, 12, 11, 12)): a.column_dimensions[c].width = w

# ---------------- Lager ----------------
l = wb.create_sheet("Lager")
l["A1"] = "Lagerliste Werkstatt"; l["A1"].font = Font(bold=True, size=14); l["F1"] = "gezählt: 09.10. (Ayşe)"
LH = ["Art.-Nr", "Artikel", "Bestand", "Einheit", "Mindestbestand", "EK-Preis €", "Wert €", "Nachbestellen", "Lieferant"]
for i, h in enumerate(LH, 1): x = l.cell(3, i, h); x.font = W; x.fill = HEAD; x.border = BOX
items = [("A-1001", "Fliesenkleber C2TE 25 kg", 14, "Sack", 10, 14.9, "Baustoff Kaiser"), ("A-1002", "Flexkleber S1 25 kg", 4, "Sack", 8, 17.5, "Baustoff Kaiser"),
         ("A-1010", "Fugenmörtel grau 5 kg", "9", "Eimer", 6, 11.4, "Weber KG"), ("A-1011", "Fugenmörtel anthrazit 5 kg", 2, "Eimer", 4, 12.3, "Weber KG"),
         ("A-1020", "Silikon sanitär weiß", 18, "Kartusche", 12, 6.2, "Baustoff Kaiser"), ("A-1021", "Silikon sanitär grau", 5, "Kartusche", 12, "6,20", "Baustoff Kaiser"),
         ("A-1030", "Dichtschlämme 2K 15 kg", 3, "Eimer", 2, 64, "Baustoff Kaiser"), ("A-1031", "Dichtband 10 m", 7, "Rolle", 5, 18.9, "Baustoff Kaiser"),
         ("A-1040", "Tiefengrund 10 l", 1, "Kanister", 3, 29.5, "Baustoff Kaiser"), ("A-1050", "Abschlussprofil Alu 2,5 m", 22, "Stk", 10, 9.8, "Weber KG"),
         ("A-1060", "Feinsteinzeug 60x60 anthrazit", "12,5", "m²", 0, 32.9, "Weber KG"), ("A-1061", "Wandfliese 30x60 weiß matt", "ca. 30", "m²", 0, 24.5, "Weber KG"),
         ("A-1070", "Fliesenkreuze 3 mm", 9, "Beutel", 3, 3.9, "Baustoff Kaiser"), ("A-1080", "Entkopplungsmatte 30 m²", 2, "Rolle", 1, 189, "Weber KG")]
for i, (nr, art, best, eh, mi, ek, lf) in enumerate(items):
    rr = 4 + i
    for c, v in enumerate((nr, art, best, eh, mi, ek), 1): l.cell(rr, c, v)
    l.cell(rr, 6).number_format = EUR
    l.cell(rr, 7, f'=IFERROR(C{rr}*F{rr},"")').number_format = EUR
    l.cell(rr, 8, f'=IF(AND(ISNUMBER(C{rr}),C{rr}<E{rr}),"JA","")')
    l.cell(rr, 9, lf)
    for c in range(1, 10): l.cell(rr, c).border = BOX
ll = 4 + len(items) - 1
l.cell(ll + 2, 2, "Lagerwert gesamt").font = B; l.cell(ll + 2, 7, f"=SUM(G4:G{ll})").number_format = EUR; l.cell(ll + 2, 7).font = B
l.conditional_formatting.add(f"H4:H{ll}", CellIsRule(operator="equal", formula=['"JA"'], fill=PatternFill("solid", fgColor="F4CCCC"), font=Font(bold=True, color="9C0006")))
for c, w in zip("ABCDEFGHI", (9, 28, 9, 10, 14, 11, 11, 13, 16)): l.column_dimensions[c].width = w

# ---------------- Übersicht: Kennzahlen ----------------
kpis = [("Stunden Oktober", f"='Stunden Oktober'!H{STD_TOTAL}", "0.00"), ("Lohnkosten Oktober", f"='Stunden Oktober'!I{STD_TOTAL}", EUR),
        ("Rechnungen offen (brutto)", f"=Aufträge!H{lr+3}", EUR), ("Umsatz netto 2026", f"=Aufträge!E{lr+2}", EUR),
        ("Lagerwert", f"=Lager!G{ll+2}", EUR), ("Artikel nachbestellen", f'=COUNTIF(Lager!H4:H{ll},"JA")', "0")]
for i, (t, f, nf) in enumerate(kpis):
    u.cell(5 + i, 2, t).font = B; c = u.cell(5 + i, 4, f); c.number_format = nf; c.font = Font(bold=True, size=13, color="0F6B5C")
u.column_dimensions["B"].width = 28; u.column_dimensions["D"].width = 16

tmp = tempfile.mkdtemp(); raw = os.path.join(tmp, os.path.basename(OUT)); wb.save(raw)
if shutil.which("soffice"):
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", os.path.join(tmp, "c"), raw], capture_output=True)
    raw = os.path.join(tmp, "c", os.path.basename(OUT))
shutil.copy(raw, OUT); print(OUT, "Stunden GESAMT in Zeile", STD_TOTAL)
