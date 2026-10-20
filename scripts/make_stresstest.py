"""Absichtlich chaotische Excel-Datei, wie sie in echten Betrieben vorkommt (Stresstest für Zettelfrei).
    python scripts/make_stresstest.py   (openpyxl; LibreOffice optional für zwischengespeicherte Formelwerte)
"""
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation

wb = Workbook()
B = Font(bold=True); RED = Font(color="C00000", bold=True)
Y = PatternFill("solid", fgColor="FFF2CC"); G = PatternFill("solid", fgColor="D9EAD3"); GR = PatternFill("solid", fgColor="D9D9D9")
thin = Side(style="thin", color="999999"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

# ---------- 1) Stunden: Titel, zweizeiliger Kopf, Zwischensummen, Leerzeilen, gemischte Formate ----------
ws = wb.active; ws.title = "Std. Okt+Nov (NEU!!)"
ws["B1"] = "Krause Fliesen & Bad GmbH"; ws["B1"].font = Font(bold=True, size=16)
ws["B2"] = "Stundenerfassung – NICHT LÖSCHEN!!! (Stand: siehe unten)"; ws["B2"].font = RED
ws["H2"] = "Stundensatz:"; ws["I2"] = 48; ws["I2"].number_format = '#,##0.00 "€"'
ws["B3"] = None
# zweizeiliger, verbundener Kopf ab Spalte B
ws.merge_cells("B4:B5"); ws["B4"] = "Datum"
ws.merge_cells("C4:D4"); ws["C4"] = "Wer / Wo"
ws["C5"] = "MA"; ws["D5"] = "Baustelle / Kunde"
ws.merge_cells("E4:G4"); ws["E4"] = "Zeit"
ws["E5"] = "von"; ws["F5"] = "bis"; ws["G5"] = "Pause\n(min)"
ws.merge_cells("H4:H5"); ws["H4"] = "Std."
ws.merge_cells("I4:I5"); ws["I4"] = "Betrag"
ws.merge_cells("J4:J5"); ws["J4"] = "Bemerkung / Material?"
ws.merge_cells("K4:K5"); ws["K4"] = "Std."   # doppelte Überschrift!
for c in "BCDEFGHIJK":
    for r in (4, 5):
        ws[f"{c}{r}"].font = B; ws[f"{c}{r}"].fill = GR; ws[f"{c}{r}"].border = BOX
        ws[f"{c}{r}"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

rows = [
    # Datum in allen Varianten, Zeiten als Text/Zahl/Uhrzeit, Pausen mal Zahl mal Text
    (dt.date(2026,10,1), "Mehmet", "Bad Schneider", "7:00", "16:30", 30, None, "Wand gefliest"),
    ("02.10.26", "mehmet ", "Bad  Schneider", dt.time(7,15), dt.time(15,45), "30 min", None, ""),
    ("2.10.2026", "Kevin", "Küche Yılmaz", 0.2916666, 0.6875, 45, None, "Silikon fehlt!!"),
    ("05.10.2026", "Jonas", "Praxis Dr. Hahn", "07:00", "17:45", 0, None, "Pause vergessen?"),
    (None, None, None, None, None, None, None, None),   # Leerzeile mitten drin
    ("06.10.2026", "Kevin + Jonas", "Hahn", "8", "16", "1/2 Std", None, "zu zweit"),
    ("Okt 07", "Mehmet", "Bad Schneider", "7:00", "15:00", 30, None, "Fugen"),
    ("08.10.2026", "M. Öztürk", "Lager / Werkstatt", "6:30", "11:00", None, None, "krank ab Mittag"),
]
r = 6
for d, ma, bs, von, bis, p, _, bem in rows:
    ws.cell(r, 2, d); ws.cell(r, 3, ma); ws.cell(r, 4, bs); ws.cell(r, 5, von); ws.cell(r, 6, bis); ws.cell(r, 7, p); ws.cell(r, 10, bem)
    if isinstance(d, dt.date): ws.cell(r, 2).number_format = "DD.MM.YYYY"
    for c in (5, 6):
        if isinstance(ws.cell(r, c).value, (dt.time, float)): ws.cell(r, c).number_format = "hh:mm"
    if von is not None:
        ws.cell(r, 8, f"=IFERROR(ROUND(((F{r}-E{r})*24)-IF(ISNUMBER(G{r}),G{r}/60,0),2),\"?\")")  # WENNFEHLER/ISTZAHL
        ws.cell(r, 9, f"=IF(ISNUMBER(H{r}),H{r}*$I$2,\"\")"); ws.cell(r, 9).number_format = '#,##0.00 "€"'
        ws.cell(r, 11, f"=H{r}")
    r += 1
# Zwischensumme mitten in der Tabelle
ws.cell(r, 2, "Zwischensumme Oktober (1. Woche)").font = B
ws.cell(r, 8, f"=SUBTOTAL(9,H6:H{r-1})"); ws.cell(r, 9, f"=SUBTOTAL(9,I6:I{r-1})"); ws.cell(r, 9).number_format = '#,##0.00 "€"'
for c in range(2, 12): ws.cell(r, c).fill = Y
sub1 = r; r += 1
ws.cell(r, 2, "--- November ---").font = Font(italic=True, color="7F7F7F"); r += 1
start2 = r
for d, ma, bs, von, bis, p, bem in [
    (dt.datetime(2026,11,2,0,0), "Mehmet", "Neubau Weber (EFH)", "7:00", "16:00", 30, "Estrich noch feucht"),
    (46329, "Kevin", "Neubau Weber (EFH)", "7:00", "16:00", 30, "Datum als Zahl"),
    ("03/11/2026", "Jonas", "Weber", "7:30", "18:30", 30, "11 Std.!"),
    ("04.11.2026", "Jonas", "Weber", "5:00", "12:00", 15, "Ruhezeit?"),
]:
    ws.cell(r, 2, d); ws.cell(r, 3, ma); ws.cell(r, 4, bs); ws.cell(r, 5, von); ws.cell(r, 6, bis); ws.cell(r, 7, p); ws.cell(r, 10, bem)
    ws.cell(r, 2).number_format = "DD.MM.YYYY"
    ws.cell(r, 8, f"=IFERROR(ROUND(((F{r}-E{r})*24)-IF(ISNUMBER(G{r}),G{r}/60,0),2),\"?\")")
    ws.cell(r, 9, f"=IF(ISNUMBER(H{r}),H{r}*$I$2,\"\")"); ws.cell(r, 9).number_format = '#,##0.00 "€"'
    ws.cell(r, 11, f"=H{r}")
    r += 1
ws.cell(r, 2, "GESAMT").font = B
ws.cell(r, 8, f"=SUBTOTAL(9,H6:H{r-1})"); ws.cell(r, 9, f"=SUBTOTAL(9,I6:I{r-1})"); ws.cell(r, 9).number_format = '#,##0.00 "€"'
for c in range(2, 12): ws.cell(r, c).fill = G
r += 2
ws.cell(r, 2, "Stand:"); ws.cell(r, 3, "=TODAY()"); ws.cell(r, 3).number_format = "DD.MM.YYYY"
ws.cell(r+1, 2, "Chef sagt: Überstunden ab 01.11. extra eintragen!!")
# Nebentabelle rechts daneben
ws["M4"] = "Mitarbeiter"; ws["N4"] = "Satz"; ws["M4"].font = B; ws["N4"].font = B
for i, (n, s) in enumerate([("Mehmet", 52), ("Kevin", 45), ("Jonas", 45), ("M. Öztürk", 38)]):
    ws.cell(5+i, 13, n); ws.cell(5+i, 14, s)
ws["E7"].comment = Comment("Uhrzeit als echte Zeit, nicht Text", "Petra")
ws.conditional_formatting.add("H6:H40", CellIsRule(operator="greaterThan", formula=["10"], fill=PatternFill("solid", fgColor="F4CCCC")))
ws.freeze_panes = "C6"; ws.auto_filter.ref = f"B5:K{r-3}"
ws.column_dimensions["L"].hidden = True; ws["L6"] = "geheim"
for c, w in zip("ABCDEFGHIJKMN", (2, 12, 12, 22, 8, 8, 8, 7, 11, 24, 6, 12, 6)): ws.column_dimensions[c].width = w

# ---------- 2) Material: Artikel-Nr mit führenden Nullen, Zahlen als Text, Einheiten im Wert ----------
m = wb.create_sheet("Lager")
m["A1"] = "LAGERLISTE"; m["A1"].font = Font(bold=True, size=14)
m["A2"] = "zuletzt gezählt: KW 41 (Kevin)"
heads = ["Art.Nr.", "Bezeichnung", "Bestand", "Einheit", "Min.", "EK netto", "Wert", "Lieferant", "Lieferant"]  # doppelt
m.append([]); m.append(heads)
for c in range(1, len(heads)+1): m.cell(4, c).font = B; m.cell(4, c).fill = GR
items = [
    ("000123", "Fliesenkleber C2TE S1 25kg", "14", "Sack", 10, "14,90 €", "Baustoff Kaiser GmbH", "Kaiser"),
    (123, "Fliesenkleber C2TE S1 25 kg (alt)", 3, "Sack", 0, 13.5, "Kaiser", ""),
    ("A-1010", "Fugenmörtel grau 5kg", "9 Eimer", "", 6, "11.40", "Weber", ""),
    ("A-1011", "Fugenmörtel anthrazit 5kg", "", "Eimer", 4, "12,30", "Weber", "nachbestellt?"),
    ("A–1020", "Silikon sanitär weiß 310ml", 18, "Kart.", 12, 6.2, "Kaiser", ""),   # Gedankenstrich statt Bindestrich
    ("A-1021 ", "Silikon sanitär grau 310ml", "5", "Kartusche", 12, "6,20", "Kaiser", ""),  # Leerzeichen am Ende
    ("A-1060", "Feinsteinzeug 60x60 anthrazit", "12,5", "m²", 0, "32,90 €/m²", "Weber", ""),
    ("A-1061", "Wandfliese 30x60 weiß matt", "ca. 30", "qm", 0, 24.5, "Weber", ""),
    ("#NV", "Dichtband", "#BEZUG!", "Rolle", 5, 18.9, "", ""),
]
for i, it in enumerate(items):
    rr = 5 + i
    for c, v in enumerate(it[:6], 1): m.cell(rr, c, v)
    m.cell(rr, 7, f"=IFERROR(C{rr}*F{rr},\"\")"); m.cell(rr, 7).number_format = '#,##0.00 "€"'
    m.cell(rr, 8, it[6]); m.cell(rr, 9, it[7])
m.cell(15, 2, "Summe Lagerwert:"); m.cell(15, 7, "=SUM(G5:G13)"); m.cell(15, 7).font = B
dv = DataValidation(type="list", formula1='"Sack,Eimer,Kartusche,m²,Rolle,Stk"'); m.add_data_validation(dv); dv.add("D5:D40")

# ---------- 3) Rechnungen: SVERWEIS auf Kundenblatt, Status per Dropdown, Fälligkeit per Formel ----------
k = wb.create_sheet("Kunden"); k.sheet_state = "hidden"
k.append(["KdNr", "Name", "Zahlungsziel (Tage)", "E-Mail"])
for row in [(1001, "Becker, Thomas", 14, "t.becker@example.de"), (1002, "Yılmaz GbR", 30, "info@yilmaz-gbr.example"),
            (1003, "Praxis Dr. Hahn", 10, "praxis@hahn.example"), (1004, "Weber Bau GmbH & Co. KG", 30, "rechnung@weber-bau.example")]:
    k.append(row)
rw = wb.create_sheet("Rechnungen 2026")
rw["A1"] = "Ausgangsrechnungen"; rw["A1"].font = Font(bold=True, size=14)
rw.append(["Re-Nr", "KdNr", "Kunde", "Datum", "Netto", "USt %", "USt", "Brutto", "Fällig", "Status", "bezahlt am"])
for c in range(1, 12): rw.cell(2, c).font = B; rw.cell(2, c).fill = GR
inv = [("2026-031", 1001, dt.date(2026,9,15), 8865.5, 0.19, "bezahlt", dt.date(2026,9,30)),
       ("2026-032", 1002, dt.date(2026,9,28), 4120, 0.19, "offen", None),
       ("2026-033", 1003, dt.date(2026,10,2), 2380, 0.07, "offen", None),      # falscher Steuersatz
       ("2026-033", 1004, dt.date(2026,10,9), "12.400,00", 0.19, "Teilzahlung", None),   # doppelte Nr, Text-Betrag
       ("2026-035", 9999, dt.date(2026,10,12), 950, 19, "OFFEN", None)]        # unbekannter Kunde, 19 statt 19 %
for i, (nr, kd, d, net, st, stat, paid) in enumerate(inv):
    rr = 3 + i
    rw.cell(rr, 1, nr); rw.cell(rr, 2, kd)
    rw.cell(rr, 3, f'=IFERROR(VLOOKUP(B{rr},Kunden!$A:$D,2,FALSE),"??? unbekannt")')
    rw.cell(rr, 4, d).number_format = "DD.MM.YYYY"; rw.cell(rr, 5, net); rw.cell(rr, 5).number_format = '#,##0.00 "€"'
    rw.cell(rr, 6, st); rw.cell(rr, 6).number_format = "0%"
    rw.cell(rr, 7, f"=IFERROR(E{rr}*F{rr},0)"); rw.cell(rr, 8, f"=IFERROR(E{rr}+G{rr},E{rr})")
    rw.cell(rr, 9, f'=D{rr}+IFERROR(VLOOKUP(B{rr},Kunden!$A:$C,3,FALSE),14)'); rw.cell(rr, 9).number_format = "DD.MM.YYYY"
    rw.cell(rr, 10, stat); rw.cell(rr, 11, paid)
    if paid: rw.cell(rr, 11).number_format = "DD.MM.YYYY"
    for c in (7, 8): rw.cell(rr, c).number_format = '#,##0.00 "€"'
rw.cell(9, 4, "offen gesamt:"); rw.cell(9, 8, '=SUMIF(J3:J7,"offen",H3:H7)'); rw.cell(9, 8).font = B
dv2 = DataValidation(type="list", formula1='"offen,bezahlt,Teilzahlung,Mahnung"'); rw.add_data_validation(dv2); dv2.add("J3:J50")

# ---------- 4) Übersicht: nur Formeln über andere Blätter ----------
u = wb.create_sheet("Übersicht", 0)
u["A1"] = "Übersicht (automatisch – nichts eintragen!)"; u["A1"].font = Font(bold=True, size=14)
u["A3"] = "Stunden gesamt"; u["B3"] = f"='Std. Okt+Nov (NEU!!)'!H{start2+4}"
u["A4"] = "Lagerwert"; u["B4"] = "=Lager!G15"; u["B4"].number_format = '#,##0.00 "€"'
u["A5"] = "Offene Rechnungen"; u["B5"] = "='Rechnungen 2026'!H9"; u["B5"].number_format = '#,##0.00 "€"'
u["A6"] = "Anzahl Rechnungen"; u["B6"] = "=COUNTA('Rechnungen 2026'!A3:A20)"

import os, subprocess, shutil, tempfile
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
out = os.path.join(ROOT, "beispiele", "Stresstest_Chaos_Tabelle.xlsx")
tmp = tempfile.mkdtemp(); raw = os.path.join(tmp, os.path.basename(out)); wb.save(raw)
# mit LibreOffice einmal durchrechnen, damit die Datei wie aus Excel zwischengespeicherte Werte hat
if shutil.which("soffice"):
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", os.path.join(tmp, "c"), raw], capture_output=True)
    raw = os.path.join(tmp, "c", os.path.basename(out))
shutil.copy(raw, out); print(out)
